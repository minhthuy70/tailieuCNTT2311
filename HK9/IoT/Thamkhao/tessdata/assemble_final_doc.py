# -*- coding: utf-8 -*-
"""
Master Assembly Script to generate the complete Word Document
Target File: d:\CNTT2311\tailieuCNTT2311\HK9\IoT\w\MUC_LUC_CHI_TIET_VA_SO_TAY_TRA_CUU_THI_DE_MO_IOT_C123.docx
"""
import sys
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from word_helpers import (
    configure_doc_styles, add_native_toc, add_heading1, add_heading2, add_callout
)
from build_part1 import build_cover_and_info, build_part1_keyword_index
from build_part2_c1c2 import build_part2_chapter1, build_part2_chapter2
from build_part3_c3 import build_part2_chapter3
from build_part4_matrices import build_part3_and_part4

def main():
    print("Initializing document...")
    doc = docx.Document()

    # Configure Heading Styles
    configure_doc_styles(doc)

    # Configure Margins and Header/Footer
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.page_width = Inches(8.27)   # A4
        section.page_height = Inches(11.69) # A4

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("MÔN HỌC: PHÁT TRIỂN ỨNG DỤNG IOT (TT188) — SỔ TAY TRA CỨU & MỤC LỤC CHI TIẾT THI ĐỀ MỞ")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 144, 156)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun_left = fp.add_run("Tài liệu dùng cho thi đề mở | Chương 1, 2, 3       —       Trang ")
        frun_left.font.name = "Calibri"
        frun_left.font.size = Pt(9)
        frun_left.font.color.rgb = RGBColor(120, 144, 156)

        # Page Number Field
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        run_p = fp.add_run()
        run_p._r.append(fldChar1)
        run_p._r.append(instrText)
        run_p._r.append(fldChar2)
        run_p._r.append(fldChar3)
        run_p.font.name = "Calibri"
        run_p.font.size = Pt(9)
        run_p.font.bold = True
        run_p.font.color.rgb = RGBColor(15, 76, 129)

    print("Building Cover & Abbreviations...")
    build_cover_and_info(doc)
    
    # Native TOC & Navigation Tips
    add_heading2(doc, "MỤC LỤC TỰ ĐỘNG & ĐIỀU HƯỚNG NHANH TRONG MICROSOFT WORD")
    add_callout(
        doc,
        "MẸO ĐIỀU HƯỚNG TRONG PHÒNG THI KHI DÙNG MÁY TÍNH",
        "• Bật thanh điều hướng nhanh: Nhấn tổ hợp phím Ctrl + F trong Microsoft Word -> Chọn tab 'Headings' (hoặc 'Tiêu đề'). Toàn bộ cây thư mục mục lục sẽ hiện ra ở cột bên trái để bạn click nhảy ngay tới mục cần tra cứu!\n"
        "• Nếu in ra giấy: Sử dụng số trang in ở góc phải dưới và Bảng chỉ mục từ khóa A-Z tại Phần I để lật trang tức thì.",
        bg="F8FAFC", border_color="0EA5E9"
    )
    add_native_toc(doc)
    doc.add_page_break()

    print("Building Part 1: Quick Keyword Index...")
    build_part1_keyword_index(doc)
    doc.add_page_break()

    print("Building Part 2: Chapter 1...")
    build_part2_chapter1(doc)
    doc.add_page_break()

    print("Building Part 2: Chapter 2...")
    build_part2_chapter2(doc)
    doc.add_page_break()

    print("Building Part 2: Chapter 3...")
    build_part2_chapter3(doc)
    doc.add_page_break()

    print("Building Part 3 & 4: Comparison Matrices & Tactics...")
    build_part3_and_part4(doc)

    output_path = r"d:\CNTT2311\tailieuCNTT2311\HK9\IoT\w\MUC_LUC_CHI_TIET_VA_SO_TAY_TRA_CUU_THI_DE_MO_IOT_C123.docx"
    doc.save(output_path)
    print(f"SUCCESS! Complete document saved to: {output_path}")

if __name__ == "__main__":
    main()
