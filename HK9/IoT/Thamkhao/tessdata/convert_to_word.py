import sys
import os
import re
import time
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from markdown_it import MarkdownIt

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def sanitize_text(text):
    """Loại bỏ ký tự điều khiển không hợp lệ trong chuẩn XML 1.0."""
    if not text:
        return ""
    return re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]', '', text)

def calc_smart_col_widths(num_cols, col_max_lens, total_width=6.5):
    """Tính toán tỷ lệ độ rộng cột thông minh cho bảng Word dựa trên độ dài dữ liệu."""
    if num_cols <= 0:
        return []
    min_w = 0.8 if num_cols <= 4 else (0.5 if num_cols <= 6 else 0.4)
    weights = [max(1.0, (max(1, l) ** 0.6)) for l in col_max_lens]
    total_weight = sum(weights)
    widths = [(w / total_weight) * total_width for w in weights]
    widths = [max(min_w, w) for w in widths]
    current_sum = sum(widths)
    return [round((w / current_sum) * total_width, 2) for w in widths]

def add_paragraph_bottom_border(paragraph, color="1E3A8A", sz="18"):
    """Thêm đường viền trang trí nằm dưới tiêu đề đoạn văn."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="6" w:color="{color}"/></w:pBdr>')
    pPr.append(pBdr)

def create_word_document(md_path, docx_path):
    print(f"Bắt đầu chuyển đổi: {md_path} -> {docx_path}")
    start_time = time.time()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = docx.Document()
    
    # 1. Thiết lập trang A4 chuẩn xuất bản học thuật
    section = doc.sections[0]
    section.page_width = Inches(8.27)    # 21.0 cm (A4)
    section.page_height = Inches(11.69)  # 29.7 cm
    section.top_margin = Inches(0.79)    # 2.0 cm
    section.bottom_margin = Inches(0.79) # 2.0 cm
    section.left_margin = Inches(0.98)   # 2.5 cm (chuẩn gáy sách đóng tập)
    section.right_margin = Inches(0.79)  # 2.0 cm
    
    # Trang bìa không hiển thị Header và Footer
    section.different_first_page_header_footer = True
    
    # Thiết lập Header (Từ trang 2 trở đi) có đường kẻ chỉ trang nhã
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_paragraph_bottom_border(header_p, color="CBD5E1", sz="6")
    hrun = header_p.add_run("GIÁO TRÌNH INTERNET OF THINGS (IoT)  |  ĐH KỸ THUẬT – CÔNG NGHỆ CẦN THƠ")
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(100, 116, 139)

    # Thiết lập Footer (Từ trang 2 trở đi)
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = footer_p.add_run("— Trang ")
    frun.font.name = "Times New Roman"
    frun.font.size = Pt(9.5)
    frun.font.color.rgb = RGBColor(100, 116, 139)
    
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    footer_p._p.append(fldSimple)

    frun2 = footer_p.add_run(" / ")
    frun2.font.name = "Times New Roman"
    frun2.font.size = Pt(9.5)
    frun2.font.color.rgb = RGBColor(100, 116, 139)

    fldSimple2 = OxmlElement('w:fldSimple')
    fldSimple2.set(qn('w:instr'), 'NUMPAGES')
    footer_p._p.append(fldSimple2)

    frun3 = footer_p.add_run(" —")
    frun3.font.name = "Times New Roman"
    frun3.font.size = Pt(9.5)
    frun3.font.color.rgb = RGBColor(100, 116, 139)

    # 2. Cấu hình Style mặc định (Normal)
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(15, 23, 42)
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(4.5)
    style_normal.paragraph_format.space_before = Pt(0)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 3. Tạo Trang Bìa Học Thuật Chính Quy Sang Trọng
    build_cover_page(doc)
    doc.add_page_break()

    # 4. Tiền xử lý chuẩn hóa cú pháp Markdown song ngữ (*...*) để luôn hiển thị in nghiêng chuẩn
    md_text = re.sub(r'\s+\*\)', '*)', md_text)
    md_text = re.sub(r'\*\s+\)', '*)', md_text)
    md_text = re.sub(r'\(\s+\*', '(*', md_text)

    # 5. Phân tích cú pháp Markdown sang Tokens AST
    md = MarkdownIt().enable('table')
    tokens = md.parse(md_text)
    print(f"Tổng số token markdown đã phân tích: {len(tokens)}")

    # 5. Duyệt và chuyển đổi các Token sang phần tử Word
    idx = 0
    total_tokens = len(tokens)
    
    current_list_type = [] # stack of 'bullet' or 'ordered'
    list_counters = []     # stack of int for ordered lists
    in_blockquote = False
    
    while idx < total_tokens:
        token = tokens[idx]
        ttype = token.type
        
        # --- TIÊU ĐỀ (HEADINGS) ---
        if ttype == 'heading_open':
            h_tag = token.tag # h1, h2, h3, h4, h5, h6
            inline_token = tokens[idx + 1]
            idx += 2 # skip heading_close
            
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            heading_text = inline_token.content.strip()
            
            if h_tag == 'h1':
                if any(k in heading_text for k in ['Chương', 'Kết luận', 'Tài liệu tham khảo', 'Thông tin xuất bản', 'Mục lục', 'Phần mở đầu']):
                    p.paragraph_format.page_break_before = True
                
                # Định dạng đặc biệt cho Tiêu đề Chương lớn
                if heading_text.startswith("Chương "):
                    parts = heading_text.split(":", 1)
                    ch_label = parts[0].strip().upper() # CHƯƠNG X
                    ch_title = parts[1].strip() if len(parts) > 1 else ""
                    
                    p.paragraph_format.space_before = Pt(18)
                    p.paragraph_format.space_after = Pt(12)
                    add_paragraph_bottom_border(p, color="1E3A8A", sz="18") # Đường kẻ xanh đậm dưới tiêu đề
                    
                    r_lbl = p.add_run(f"{ch_label}\n")
                    r_lbl.font.name = "Times New Roman"
                    r_lbl.font.size = Pt(13)
                    r_lbl.font.bold = True
                    r_lbl.font.color.rgb = RGBColor(37, 99, 235) # Xanh dương nổi bật
                    
                    r_ttl = p.add_run(ch_title.upper())
                    r_ttl.font.name = "Times New Roman"
                    r_ttl.font.size = Pt(19)
                    r_ttl.font.bold = True
                    r_ttl.font.color.rgb = RGBColor(30, 58, 138) # Xanh Navy sâu
                else:
                    p.paragraph_format.space_before = Pt(20)
                    p.paragraph_format.space_after = Pt(10)
                    add_paragraph_bottom_border(p, color="1E3A8A", sz="16")
                    render_inline(p, inline_token.children, default_size=Pt(17), default_bold=True, 
                                  default_color=RGBColor(30, 58, 138))
                    
            elif h_tag == 'h2':
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(4.5)
                render_inline(p, inline_token.children, default_size=Pt(14), default_bold=True, 
                              default_color=RGBColor(30, 58, 138))
            elif h_tag == 'h3':
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(3.5)
                render_inline(p, inline_token.children, default_size=Pt(12.5), default_bold=True, 
                              default_color=RGBColor(29, 78, 216))
            elif h_tag == 'h4':
                p.paragraph_format.space_before = Pt(7)
                p.paragraph_format.space_after = Pt(2.5)
                render_inline(p, inline_token.children, default_size=Pt(11.5), default_bold=True, 
                              default_color=RGBColor(51, 65, 85))
            else: # h5, h6
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(2)
                render_inline(p, inline_token.children, default_size=Pt(11), default_bold=True, 
                              default_italic=True, default_color=RGBColor(71, 85, 105))
            
            idx += 1
            continue

        # --- ĐOẠN VĂN (PARAGRAPHS) ---
        elif ttype == 'paragraph_open':
            inline_token = tokens[idx + 1]
            idx += 2 # skip paragraph_close
            p_text = inline_token.content.strip()
            
            # 1. Phát hiện và căn giữa Chú thích hình ảnh (Figure Caption)
            if re.match(r'^\*?\*?Hình\s+[\d\.]+', p_text):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(12)
                render_inline(p, inline_token.children, default_size=Pt(10.5), default_italic=True,
                              default_color=RGBColor(71, 85, 105))
                idx += 1
                continue

            # 2. Phát hiện Tiêu đề bảng (Table Caption)
            elif re.match(r'^\*?\*?Bảng\s+[\d\.]+', p_text):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                render_inline(p, inline_token.children, default_size=Pt(10.5), default_bold=True,
                              default_color=RGBColor(30, 58, 138))
                idx += 1
                continue

            # 3. Nằm trong danh sách
            elif current_list_type:
                p = doc.add_paragraph()
                level = len(current_list_type)
                p.paragraph_format.left_indent = Inches(0.25 * level)
                p.paragraph_format.first_line_indent = -Inches(0.18)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                
                if current_list_type[-1] == 'bullet':
                    bullet_sym = "•  " if level == 1 else ("–  " if level == 2 else "▪  ")
                    brun = p.add_run(bullet_sym)
                    brun.font.name = "Times New Roman"
                    brun.font.bold = True
                    brun.font.color.rgb = RGBColor(37, 99, 235)
                else: # ordered
                    list_counters[-1] += 1
                    num_sym = f"{list_counters[-1]}. "
                    nrun = p.add_run(num_sym)
                    nrun.font.name = "Times New Roman"
                    nrun.font.bold = True
                    nrun.font.color.rgb = RGBColor(30, 58, 138)
                
                render_inline(p, inline_token.children)
            elif in_blockquote:
                add_callout_box(doc, inline_token.children)
            else:
                p = doc.add_paragraph()
                render_inline(p, inline_token.children)
                
            idx += 1
            continue

        # --- DANH SÁCH (LISTS) ---
        elif ttype == 'bullet_list_open':
            current_list_type.append('bullet')
            idx += 1
            continue
        elif ttype == 'bullet_list_close':
            if current_list_type:
                current_list_type.pop()
            idx += 1
            continue
        elif ttype == 'ordered_list_open':
            current_list_type.append('ordered')
            list_counters.append(0)
            idx += 1
            continue
        elif ttype == 'ordered_list_close':
            if current_list_type:
                current_list_type.pop()
            if list_counters:
                list_counters.pop()
            idx += 1
            continue
        elif ttype == 'list_item_open' or ttype == 'list_item_close':
            idx += 1
            continue

        # --- KHỐI MÃ NGUỒN (CODE BLOCKS / FENCES) ---
        elif ttype == 'fence':
            add_code_block(doc, token.content, token.info)
            idx += 1
            continue

        # --- BẢNG DỮ LIỆU (TABLES) ---
        elif ttype == 'table_open':
            table_tokens = []
            while idx < total_tokens and tokens[idx].type != 'table_close':
                table_tokens.append(tokens[idx])
                idx += 1
            if idx < total_tokens:
                table_tokens.append(tokens[idx])
            
            build_docx_table(doc, table_tokens)
            idx += 1
            continue

        # --- TRÍCH DẪN (BLOCKQUOTE) ---
        elif ttype == 'blockquote_open':
            in_blockquote = True
            idx += 1
            continue
        elif ttype == 'blockquote_close':
            in_blockquote = False
            idx += 1
            continue

        # --- ĐƯỜNG PHÂN CÁCH (HORIZONTAL RULE) ---
        elif ttype == 'hr':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("❖  ❖  ❖")
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(148, 163, 184)
            idx += 1
            continue

        else:
            idx += 1

    # Lưu tệp Word hoàn chỉnh
    print(f"Đang ghi tệp Word ra đĩa: {docx_path} ...")
    doc.save(docx_path)
    elapsed = time.time() - start_time
    file_size_kb = os.path.getsize(docx_path) / 1024
    print(f"Hoàn thành xuất sắc trong {elapsed:.2f} giây! Dung lượng tệp: {file_size_kb:.1f} KB.")

# --- CÁC HÀM XỬ LÝ NỘI DUNG VÀ ĐỊNH DẠNG ---

def render_inline(paragraph, children, default_size=Pt(12), default_bold=False, 
                  default_italic=False, default_color=None):
    """Duyệt AST con của thẻ inline và thêm các Run với định dạng chuẩn xác."""
    if not children:
        return
    
    bold_stack = [default_bold]
    italic_stack = [default_italic]
    code_stack = [False]
    link_stack = [None]
    
    for c in children:
        ctype = c.type
        
        if ctype == 'text':
            content = sanitize_text(c.content)
            if not content:
                continue
            run = paragraph.add_run(content)
            run.font.name = 'Times New Roman'
            run.font.size = default_size
            run.bold = bold_stack[-1]
            run.italic = italic_stack[-1]
            
            if link_stack[-1]:
                run.font.color.rgb = RGBColor(37, 99, 235)
                run.underline = True
            elif default_color:
                run.font.color.rgb = default_color
            else:
                # Phân tách màu văn bản song ngữ:
                # Bản dịch tiếng Việt trong ngoặc (*...*) in nghiêng màu xám than sang trọng
                if italic_stack[-1]:
                    run.font.color.rgb = RGBColor(51, 65, 85) # Slate dark
                else:
                    run.font.color.rgb = RGBColor(15, 23, 42) # Near black
                    
        elif ctype == 'code_inline':
            run = paragraph.add_run(sanitize_text(c.content))
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.bold = False
            run.font.color.rgb = RGBColor(185, 28, 28) # Đỏ mận tinh tế cho inline code
            
        elif ctype == 'strong_open':
            bold_stack.append(True)
        elif ctype == 'strong_close':
            if len(bold_stack) > 1:
                bold_stack.pop()
                
        elif ctype == 'em_open':
            italic_stack.append(True)
        elif ctype == 'em_close':
            if len(italic_stack) > 1:
                italic_stack.pop()
                
        elif ctype == 'link_open':
            href = None
            if c.attrs:
                href = dict(c.attrs).get('href')
            link_stack.append(href or "url")
        elif ctype == 'link_close':
            if len(link_stack) > 1:
                link_stack.pop()
                
        elif ctype == 'softbreak':
            run = paragraph.add_run(' ')
            run.font.name = 'Times New Roman'
            run.font.size = default_size
            
        elif ctype == 'hardbreak':
            run = paragraph.add_run('\n')
            run.font.name = 'Times New Roman'
            run.font.size = default_size

def add_code_block(doc, code_text, lang=""):
    """Tạo hộp mã nguồn chuyên nghiệp với viền xanh dương và nhãn ngôn ngữ."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:left w:val="single" w:sz="28" w:space="0" w:color="2563EB"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>'))
    
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="120" w:type="dxa"/>'
        f'<w:bottom w:w="120" w:type="dxa"/>'
        f'<w:left w:w="160" w:type="dxa"/>'
        f'<w:right w:w="160" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    clean_code = sanitize_text(code_text.rstrip())
    if lang:
        lrun = p.add_run(f"// [MÃ NGUỒN: {lang.upper()}]\n")
        lrun.font.name = "Consolas"
        lrun.font.size = Pt(8.5)
        lrun.font.bold = True
        lrun.font.color.rgb = RGBColor(100, 116, 139)
        
    crun = p.add_run(clean_code)
    crun.font.name = "Consolas"
    crun.font.size = Pt(9.5)
    crun.font.color.rgb = RGBColor(15, 23, 42)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(4)

def add_callout_box(doc, children):
    """Tạo hộp trích dẫn ghi chú Callout với viền xanh ngọc / lam."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:left w:val="single" w:sz="28" w:space="0" w:color="0D9488"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0FDFA"/>'))
    
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="100" w:type="dxa"/>'
        f'<w:bottom w:w="100" w:type="dxa"/>'
        f'<w:left w:w="160" w:type="dxa"/>'
        f'<w:right w:w="160" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    render_inline(p, children, default_size=Pt(11), default_italic=True, default_color=RGBColor(19, 78, 74))
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def build_docx_table(doc, table_tokens):
    """Tạo bảng dữ liệu Word chuẩn xuất bản với căn chỉnh thông minh."""
    header_cells = []
    data_rows = []
    
    current_row = []
    in_th = False
    in_td = False
    
    for t in table_tokens:
        if t.type == 'th_open':
            in_th = True
        elif t.type == 'th_close':
            in_th = False
        elif t.type == 'td_open':
            in_td = True
        elif t.type == 'td_close':
            in_td = False
        elif t.type == 'inline':
            if in_th:
                header_cells.append(t)
            elif in_td:
                current_row.append(t)
        elif t.type == 'tr_close':
            if current_row:
                data_rows.append(current_row)
                current_row = []
                
    num_cols = len(header_cells)
    if num_cols == 0:
        return
        
    num_rows = len(data_rows) + 1
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Thiết lập đường viền bảng mảnh thanh nhã
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    col_max_lens = [len(h.content) for h in header_cells]
    for row in data_rows:
        for c_idx, cell in enumerate(row):
            if c_idx < len(col_max_lens):
                col_max_lens[c_idx] = max(col_max_lens[c_idx], len(cell.content))
                
    col_widths = calc_smart_col_widths(num_cols, col_max_lens, total_width=6.5)

    # 1. Hàng tiêu đề (Header Row)
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    header_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for col_idx, cell_inline in enumerate(header_cells):
        cell = table.cell(0, col_idx)
        cell.width = Inches(col_widths[col_idx])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E3A8A"/>')) # Xanh Navy đậm sang trọng
        
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="120" w:type="dxa"/>'
            f'<w:bottom w:w="120" w:type="dxa"/>'
            f'<w:left w:w="140" w:type="dxa"/>'
            f'<w:right w:w="140" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        render_inline(p, cell_inline.children, default_size=Pt(10), default_bold=True, 
                      default_color=RGBColor(255, 255, 255))

    # 2. Các hàng dữ liệu (Data Rows)
    for r_idx, row_data in enumerate(data_rows):
        row_num = r_idx + 1
        row = table.rows[row_num]
        row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        
        for c_idx, cell_inline in enumerate(row_data):
            if c_idx >= num_cols:
                break
            cell = table.cell(row_num, c_idx)
            cell.width = Inches(col_widths[c_idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            tcPr = cell._tc.get_or_add_tcPr()
            if bg_color != "FFFFFF":
                tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>'))
                
            tcMar = parse_xml(
                f'<w:tcMar {nsdecls("w")}>'
                f'<w:top w:w="90" w:type="dxa"/>'
                f'<w:bottom w:w="90" w:type="dxa"/>'
                f'<w:left w:w="120" w:type="dxa"/>'
                f'<w:right w:w="120" w:type="dxa"/>'
                f'</w:tcMar>'
            )
            tcPr.append(tcMar)
            
            p = cell.paragraphs[0]
            # Căn giữa nếu nội dung ngắn (STT, ký hiệu, bit), căn trái nếu nội dung văn bản dài
            if len(cell_inline.content) <= 12 or cell_inline.content.isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            render_inline(p, cell_inline.children, default_size=Pt(10), default_color=RGBColor(15, 23, 42))

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(6)

def build_cover_page(doc):
    """Tạo trang bìa phong cách học thuật chính quy chuẩn mực cho giáo trình."""
    # Cơ quan chủ quản & Tên trường
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_before = Pt(10)
    p_top.paragraph_format.space_after = Pt(2)
    
    r_sub = p_top.add_run("ỦY BAN NHÂN DÂN THÀNH PHỐ CẦN THƠ\n")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(11.5)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    
    r_univ = p_top.add_run("TRƯỜNG ĐẠI HỌC KỸ THUẬT – CÔNG NGHỆ CẦN THƠ\n")
    r_univ.font.name = "Times New Roman"
    r_univ.font.size = Pt(13.5)
    r_univ.font.bold = True
    r_univ.font.color.rgb = RGBColor(30, 58, 138)
    
    r_dept = p_top.add_run("KHOA CÔNG NGHỆ THÔNG TIN – BỘ MÔN KỸ THUẬT PHẦN MỀM & IoT")
    r_dept.font.name = "Times New Roman"
    r_dept.font.size = Pt(11.5)
    r_dept.font.bold = True
    r_dept.font.color.rgb = RGBColor(37, 99, 235)

    # Đường kẻ phân cách
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(60)
    r_line = p_div.add_run("━━━━━━━━━━━━━━━━━━━")
    r_line.font.name = "Times New Roman"
    r_line.font.size = Pt(12)
    r_line.font.color.rgb = RGBColor(203, 213, 225)

    # Tiêu đề giáo trình lớn
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(12)
    
    r_type = p_title.add_run("GIÁO TRÌNH\n")
    r_type.font.name = "Times New Roman"
    r_type.font.size = Pt(15)
    r_type.font.bold = True
    r_type.font.color.rgb = RGBColor(71, 85, 105)

    r_main = p_title.add_run("INTERNET OF THINGS (IoT)\n")
    r_main.font.name = "Times New Roman"
    r_main.font.size = Pt(28)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(30, 58, 138)

    r_subt = p_title.add_run("Kiến Trúc, Giao Thức Mạng, Phần Cứng Nhúng & Ứng Dụng Thực Tế")
    r_subt.font.name = "Times New Roman"
    r_subt.font.size = Pt(13.5)
    r_subt.font.italic = True
    r_subt.font.color.rgb = RGBColor(37, 99, 235)

    # Họa tiết kim cương
    p_orn = doc.add_paragraph()
    p_orn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_orn.paragraph_format.space_before = Pt(18)
    p_orn.paragraph_format.space_after = Pt(40)
    r_orn = p_orn.add_run("❖  ❖  ❖")
    r_orn.font.name = "Times New Roman"
    r_orn.font.size = Pt(11)
    r_orn.font.color.rgb = RGBColor(148, 163, 184)

    # Tác giả & Giảng viên
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_before = Pt(20)
    p_author.paragraph_format.space_after = Pt(4)
    
    r_auth_title = p_author.add_run("Chủ biên: ")
    r_auth_title.font.name = "Times New Roman"
    r_auth_title.font.size = Pt(12.5)
    r_auth_title.font.italic = True
    
    r_auth_name = p_author.add_run("TS. NGUYỄN ĐÌNH TỨ\n")
    r_auth_name.font.name = "Times New Roman"
    r_auth_name.font.size = Pt(13)
    r_auth_name.font.bold = True
    r_auth_name.font.color.rgb = RGBColor(15, 23, 42)

    r_unit = p_author.add_run("(Trưởng Bộ môn Cơ điện tử – Khoa Kỹ thuật Cơ khí)\n\n")
    r_unit.font.name = "Times New Roman"
    r_unit.font.size = Pt(10.5)
    r_unit.font.italic = True
    r_unit.font.color.rgb = RGBColor(100, 116, 139)
    
    r_co = p_author.add_run("Tham gia biên soạn: ")
    r_co.font.name = "Times New Roman"
    r_co.font.size = Pt(11.5)
    r_co.font.italic = True
    
    r_co_name = p_author.add_run("ThS. HUỲNH BẢO TOÀN")
    r_co_name.font.name = "Times New Roman"
    r_co_name.font.size = Pt(12)
    r_co_name.font.bold = True
    r_co_name.font.color.rgb = RGBColor(15, 23, 42)

    # Địa danh & Năm xuất bản
    p_bot = doc.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bot.paragraph_format.space_before = Pt(60)
    p_bot.paragraph_format.space_after = Pt(0)
    
    r_bot = p_bot.add_run("CẦN THƠ – NĂM 2024 / 2025")
    r_bot.font.name = "Times New Roman"
    r_bot.font.size = Pt(11.5)
    r_bot.font.bold = True
    r_bot.font.color.rgb = RGBColor(100, 116, 139)

if __name__ == '__main__':
    src = r'd:\CNTT2311\tailieuCNTT2311\HK9\IoT\IoTmain.md'
    dst = r'd:\CNTT2311\tailieuCNTT2311\HK9\IoT\IoTmain.docx'
    create_word_document(src, dst)
