# -*- coding: utf-8 -*-
"""
Helper functions for rich Word document generation with native Word Outline Styles
"""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def configure_doc_styles(doc):
    # Configure Heading 1
    s1 = doc.styles['Heading 1']
    s1.font.name = "Calibri"
    s1.font.size = Pt(16)
    s1.font.bold = True
    s1.font.color.rgb = RGBColor(15, 76, 129) # Navy
    
    # Configure Heading 2
    s2 = doc.styles['Heading 2']
    s2.font.name = "Calibri"
    s2.font.size = Pt(13)
    s2.font.bold = True
    s2.font.color.rgb = RGBColor(30, 58, 138) # Dark Blue
    
    # Configure Heading 3
    s3 = doc.styles['Heading 3']
    s3.font.name = "Calibri"
    s3.font.size = Pt(11.5)
    s3.font.bold = True
    s3.font.color.rgb = RGBColor(3, 105, 161) # Steel Blue

    # Configure Heading 4
    s4 = doc.styles['Heading 4']
    s4.font.name = "Calibri"
    s4.font.size = Pt(10.5)
    s4.font.bold = True
    s4.font.italic = True
    s4.font.color.rgb = RGBColor(51, 65, 85)

def set_cell_bg(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_table_borders(table, color="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_heading1(doc, text):
    h = doc.add_paragraph(style='Heading 1')
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 76, 129) # Navy
    return h

def add_heading2(doc, text):
    h = doc.add_paragraph(style='Heading 2')
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # Dark Blue
    return h

def add_heading3(doc, text):
    h = doc.add_paragraph(style='Heading 3')
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(3, 105, 161) # Steel Blue
    return h

def add_heading4(doc, text):
    h = doc.add_paragraph(style='Heading 4')
    h.paragraph_format.space_before = Pt(6)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(51, 65, 85)
    return h

def add_bullet(doc, text, bold_prefix=None, indent=0.25):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.15
    
    # bullet icon
    b_run = p.add_run("•  ")
    b_run.font.name = "Calibri"
    b_run.font.bold = True
    b_run.font.color.rgb = RGBColor(2, 132, 199)
    
    if bold_prefix:
        bp = p.add_run(bold_prefix)
        bp.font.name = "Calibri"
        bp.font.size = Pt(10)
        bp.font.bold = True
        bp.font.color.rgb = RGBColor(15, 23, 42)
        
    t = p.add_run(text)
    t.font.name = "Calibri"
    t.font.size = Pt(10)
    t.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_subbullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.12
    
    b_run = p.add_run("–  ")
    b_run.font.name = "Calibri"
    b_run.font.bold = True
    b_run.font.color.rgb = RGBColor(100, 116, 139)
    
    if bold_prefix:
        bp = p.add_run(bold_prefix)
        bp.font.name = "Calibri"
        bp.font.size = Pt(9.5)
        bp.font.bold = True
        bp.font.color.rgb = RGBColor(15, 23, 42)
        
    t = p.add_run(text)
    t.font.name = "Calibri"
    t.font.size = Pt(9.5)
    t.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout(doc, title, text, bg="F0F9FF", border_color="0284C7"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.77)
    set_cell_bg(cell, bg)
    set_cell_margins(cell, top=100, bottom=100, left=180, right=140)
    
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    table._tbl.tblPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    if title:
        r_title = p.add_run(f"📌 {title}\n")
        r_title.font.name = "Calibri"
        r_title.font.size = Pt(10)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(15, 76, 129)
        
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "CBD5E1")
    
    hdr_row = table.rows[0]
    hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for i, title in enumerate(headers):
        cell = hdr_row.cells[i]
        if col_widths and i < len(col_widths):
            cell.width = Inches(col_widths[i])
        set_cell_bg(cell, "1E3A8A")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            cell = row.cells[c_idx]
            if col_widths and c_idx < len(col_widths):
                cell.width = Inches(col_widths[c_idx])
            set_cell_bg(cell, bg_color)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.12
            
            if c_idx == 0 and len(cell_value) <= 6:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif c_idx in [1, 2] and len(cell_value) <= 12 and not ' ' in cell_value:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            run = p.add_run(cell_value)
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(30, 41, 59)
            
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)
    return table

def add_native_toc(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run()
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr = parse_xml(r'<w:instrText %s xml:space="preserve"> TOC \o "1-3" \h \z \u </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r._r.append(fld1)
    r._r.append(instr)
    r._r.append(fld2)
    r._r.append(fld3)

print("Updated word_helpers with native outline styles!")
